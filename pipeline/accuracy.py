"""
The accuracy sample: a blind random draw of the attributed population.

The eval suite (pipeline/evaluation.py, pipeline/targets.py) is a
regression tripwire over chosen cases; its pass rate describes the
suite. This module draws rows uniformly at random from the population
the rankings are computed over, hands them to the owner as a workbook
with no prediction in sight, reads the verdicts back, and scores the
classifier against them. It is the only producer of an accuracy figure.
"""

import logging
from collections.abc import Mapping
from pathlib import Path

import polars as pl
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from pipeline.lineage import check_config_stamps
from pipeline.schemas import (
    ACCURACY_SAMPLE_SCHEMA,
    AccuracyFigures,
    ClassAgreement,
    validate_schema,
)
from pipeline.stage import classifier_stamp_keys
from utils.constants import (
    ACCURACY_CLASS_MIN_N,
    ACCURACY_SAMPLE_N,
    ACCURACY_SAMPLE_SEED,
)

logger = logging.getLogger(__name__)

SENTIMENT_VERDICTS = ("neg", "neu", "pos")
TARGET_NONE = "none"  # the sentiment is not aimed at a player
TARGET_OTHER = "other"  # aimed at a player the row's list does not offer
REJECT_REASONS = ("alias_false_positive", "unreadable")

WORKBOOK_FILENAME = "accuracy_sample.xlsx"
SAMPLE_FILENAME = "accuracy_sample.parquet"
LABELS_SHEET = "labels"
CLASSIFIER_SHEET = "classifier"
LISTS_SHEET = "lists"
META_SHEET = "meta"
LABEL_COLUMNS = (
    "comment_id",
    "permalink",
    "post_title",
    "body",
    "sentiment",
    "target",
    "reject",
    "note",
)
CLASSIFIER_COLUMNS = (
    "comment_id",
    "sentiment",
    "confidence",
    "sentiment_player",
    "attributed_player",
    "mentioned_players",
)
COLUMN_WIDTHS = {
    "comment_id": 12,
    "permalink": 10,
    "post_title": 40,
    "body": 90,
    "sentiment": 11,
    "target": 26,
    "reject": 22,
    "note": 40,
}
MENTIONED_SEPARATOR = "|"
PERMALINK = "https://www.reddit.com/r/nba/comments/{}/comment/{}/"
SENTIMENT_STAMP_KEYS = classifier_stamp_keys("sentiment")
RATE_DECIMALS = 4
MAX_REPORTED_PROBLEMS = 40

SAMPLE_COLUMNS = (
    "comment_id",
    "permalink",
    "post_title",
    "body",
    "mentioned_players",
    "sentiment",
    "confidence",
    "sentiment_player",
    "attributed_player",
)


def draw_sample(
    df: pl.DataFrame,
    posts: pl.DataFrame | None,
    n: int = ACCURACY_SAMPLE_N,
    seed: int = ACCURACY_SAMPLE_SEED,
) -> pl.DataFrame:
    """
    Draw the sample: n attributed rows, uniformly at random, under a seed.

    The population is the fact's attributed rows with a real label, the
    set every negative rate is computed over. The draw is reproducible
    from the seed as long as the fact's row order is; the ids are what
    the workbook and the parquet carry, so a redraw is never needed to
    score. Each row gets its Reddit permalink and, when the bridge is
    given, its post title.

    Args:
        df: The fact (SENTIMENT_SCHEMA columns; error rows may be present).
        posts: The post bridge (post_id, title), or None for no titles.
        n: Rows to draw.
        seed: Sampling seed.

    Returns:
        A frame of SAMPLE_COLUMNS.

    Raises:
        ValueError: If the population holds fewer than n rows.
    """
    population = df.filter(
        pl.col("attributed_player").is_not_null() & (pl.col("sentiment") != "error")
    )
    if population.height < n:
        raise ValueError(
            f"attributed population holds {population.height:,} rows, fewer "
            f"than the {n:,} requested"
        )
    sample = population.sample(n=n, seed=seed).select(
        "comment_id",
        "link_id",
        "body",
        "mentioned_players",
        "sentiment",
        "confidence",
        "sentiment_player",
        "attributed_player",
    )
    sample = sample.with_columns(
        pl.format(
            PERMALINK,
            pl.col("link_id").str.replace("^t3_", ""),
            pl.col("comment_id"),
        ).alias("permalink")
    )
    if posts is None:
        sample = sample.with_columns(pl.lit(None, dtype=pl.String).alias("post_title"))
    else:
        titles = posts.select(
            pl.col("post_id").alias("link_id"), pl.col("title").alias("post_title")
        )
        sample = sample.join(titles, on="link_id", how="left", maintain_order="left")

    mix = {
        label: sample.filter(pl.col("sentiment") == label).height
        for label in SENTIMENT_VERDICTS
    }
    logger.info(
        f"drew {sample.height:,} of {population.height:,} attributed rows "
        f"(seed {seed}); predicted mix "
        + ", ".join(f"{label} {count:,}" for label, count in mix.items())
    )
    thin = [label for label, count in mix.items() if count < ACCURACY_CLASS_MIN_N]
    if thin:
        logger.warning(
            f"predicted class(es) {thin} fall under {ACCURACY_CLASS_MIN_N} rows in "
            f"the draw: per-class figures will be wide; consider a stratified draw"
        )
    return sample.select(*SAMPLE_COLUMNS)


def target_options(mentioned_players: list[str]) -> tuple[str, ...]:
    """
    The target list a row offers: its mentioned players, then none and other.

    Args:
        mentioned_players: The row's tracked mentions, canonical names.

    Returns:
        The dropdown's values in order.
    """
    return (*mentioned_players, TARGET_NONE, TARGET_OTHER)


def write_workbook(sample: pl.DataFrame, path: Path, stamps: Mapping[str, str]) -> None:
    """
    Write the labeling workbook.

    The labels sheet shows the comment and nothing the classifier said:
    id, permalink, post title, body, then three list-validated verdict
    columns and a note. The target list is per row (that comment's
    mentioned players, none, other), served from a hidden lists sheet.
    The classifier's side and the draw's stamps sit on hidden sheets
    for the import to join back by comment id.

    Args:
        sample: Frame from draw_sample.
        path: Destination .xlsx.
        stamps: Draw metadata (config, classifier, seed, n, drawn_at).
    """
    workbook = Workbook()
    labels = workbook.active
    labels.title = LABELS_SHEET
    labels.append(list(LABEL_COLUMNS))
    for cell in labels[1]:
        cell.font = Font(bold=True)
    labels.freeze_panes = "A2"
    for index, name in enumerate(LABEL_COLUMNS, start=1):
        labels.column_dimensions[get_column_letter(index)].width = COLUMN_WIDTHS[name]

    lists = workbook.create_sheet(LISTS_SHEET)
    lists.sheet_state = "hidden"
    classifier = workbook.create_sheet(CLASSIFIER_SHEET)
    classifier.sheet_state = "hidden"
    classifier.append(list(CLASSIFIER_COLUMNS))
    meta = workbook.create_sheet(META_SHEET)
    meta.sheet_state = "hidden"
    for key, value in stamps.items():
        meta.append([key, value])

    last_row = sample.height + 1
    sentiment_validation = DataValidation(
        type="list", formula1=f'"{",".join(SENTIMENT_VERDICTS)}"', allow_blank=True
    )
    reject_validation = DataValidation(
        type="list", formula1=f'"{",".join(REJECT_REASONS)}"', allow_blank=True
    )
    labels.add_data_validation(sentiment_validation)
    labels.add_data_validation(reject_validation)
    sentiment_validation.add(f"E2:E{last_row}")
    reject_validation.add(f"G2:G{last_row}")

    list_validations: dict[tuple[str, ...], DataValidation] = {}
    wrap = Alignment(wrap_text=True, vertical="top")
    for row_number, row in enumerate(sample.iter_rows(named=True), start=2):
        options = target_options(row["mentioned_players"])
        validation = list_validations.get(options)
        if validation is None:
            lists.append(list(options))
            list_row = len(list_validations) + 1
            last_column = get_column_letter(len(options))
            validation = DataValidation(
                type="list",
                formula1=f"{LISTS_SHEET}!$A${list_row}:${last_column}${list_row}",
                allow_blank=True,
            )
            labels.add_data_validation(validation)
            list_validations[options] = validation
        validation.add(f"F{row_number}")

        labels.append([row["comment_id"], "open", row["post_title"], row["body"]])
        labels.cell(row=row_number, column=1).number_format = "@"
        link = labels.cell(row=row_number, column=2)
        link.hyperlink = row["permalink"]
        link.style = "Hyperlink"
        labels.cell(row=row_number, column=3).alignment = wrap
        labels.cell(row=row_number, column=4).alignment = wrap

        classifier.append(
            [
                row["comment_id"],
                row["sentiment"],
                row["confidence"],
                row["sentiment_player"],
                row["attributed_player"],
                MENTIONED_SEPARATOR.join(row["mentioned_players"]),
            ]
        )
    workbook.save(path)


def _text(value: object) -> str | None:
    """A cell as stripped text; None for an empty cell."""
    if value is None:
        return None
    text = str(value).strip()
    return text or None


def read_workbook(path: Path) -> tuple[pl.DataFrame, dict[str, str]]:
    """
    Read the labeled workbook back as the accuracy sample.

    Refuses the file rather than scoring a partial or off-list set:
    every drawn id must appear exactly once, an unrejected row needs a
    sentiment and a target from its own list, a rejected row needs a
    known reason. Problems are reported by sheet row.

    Args:
        path: The .xlsx written by write_workbook and filled in.

    Returns:
        (frame under ACCURACY_SAMPLE_SCHEMA in draw order; the stamps
        from the meta sheet).

    Raises:
        ValueError: If a sheet or header is missing, or any row fails
            the checks; the message lists the rows.
    """
    workbook = load_workbook(path, read_only=True, data_only=True)
    for sheet in (LABELS_SHEET, CLASSIFIER_SHEET, META_SHEET):
        if sheet not in workbook.sheetnames:
            raise ValueError(f"{path}: sheet {sheet!r} is missing")

    stamps = {
        str(key): str(value)
        for key, value in workbook[META_SHEET].iter_rows(values_only=True)
        if key is not None and value is not None
    }

    classifier_rows = workbook[CLASSIFIER_SHEET].iter_rows(values_only=True)
    if tuple(next(classifier_rows, ())) != CLASSIFIER_COLUMNS:
        raise ValueError(f"{path}: the {CLASSIFIER_SHEET} sheet's header was changed")
    drawn: dict[str, dict] = {}
    for values in classifier_rows:
        row = dict(zip(CLASSIFIER_COLUMNS, values))
        mentioned = _text(row["mentioned_players"])
        row["mentioned_players"] = (
            mentioned.split(MENTIONED_SEPARATOR) if mentioned else []
        )
        drawn[str(row["comment_id"])] = row

    label_rows = workbook[LABELS_SHEET].iter_rows(values_only=True)
    if tuple(next(label_rows, ())) != LABEL_COLUMNS:
        raise ValueError(f"{path}: the {LABELS_SHEET} sheet's header was changed")

    problems: list[str] = []
    labels: dict[str, dict] = {}
    for row_number, values in enumerate(label_rows, start=2):
        row = {name: _text(value) for name, value in zip(LABEL_COLUMNS, values)}
        comment_id = row["comment_id"]
        if comment_id is None:
            if any(row[name] for name in ("sentiment", "target", "reject", "note")):
                problems.append(f"row {row_number}: verdicts without a comment_id")
            continue
        if comment_id not in drawn:
            problems.append(f"row {row_number}: {comment_id} is not in the draw")
            continue
        if comment_id in labels:
            problems.append(f"row {row_number}: {comment_id} appears twice")
            continue

        reject = row["reject"]
        if row["sentiment"] is None and row["target"] is None and reject is None:
            problems.append(f"row {row_number}: unlabeled")
        elif reject is not None:
            if reject not in REJECT_REASONS:
                problems.append(
                    f"row {row_number}: reject {reject!r} is not one of "
                    f"{list(REJECT_REASONS)}"
                )
        else:
            sentiment = row["sentiment"]
            if sentiment not in SENTIMENT_VERDICTS:
                problems.append(
                    f"row {row_number}: sentiment {sentiment!r} is not one of "
                    f"{list(SENTIMENT_VERDICTS)}"
                )
            options = target_options(drawn[comment_id]["mentioned_players"])
            if row["target"] not in options:
                problems.append(
                    f"row {row_number}: target {row['target']!r} is not one of "
                    f"{list(options)}"
                )
        labels[comment_id] = row

    missing = [comment_id for comment_id in drawn if comment_id not in labels]
    if missing:
        problems.append(
            f"{len(missing)} drawn row(s) are absent from the {LABELS_SHEET} sheet: "
            + ", ".join(missing[:10])
            + (", ..." if len(missing) > 10 else "")
        )
    if problems:
        shown = problems[:MAX_REPORTED_PROBLEMS]
        more = len(problems) - len(shown)
        raise ValueError(
            f"{path}: {len(problems)} problem(s) in the {LABELS_SHEET} sheet:\n  "
            + "\n  ".join(shown)
            + (f"\n  ... and {more} more" if more else "")
        )

    records = []
    for comment_id, prediction in drawn.items():
        label = labels[comment_id]
        rejected = label["reject"] is not None
        records.append(
            {
                "comment_id": comment_id,
                "sentiment": prediction["sentiment"],
                "confidence": float(prediction["confidence"]),
                "sentiment_player": _text(prediction["sentiment_player"]),
                "attributed_player": prediction["attributed_player"],
                "label_sentiment": None if rejected else label["sentiment"],
                "label_target": None if rejected else label["target"],
                "reject": label["reject"],
                "note": label["note"],
            }
        )
    frame = pl.DataFrame(records, schema=ACCURACY_SAMPLE_SCHEMA)
    validate_schema(frame, ACCURACY_SAMPLE_SCHEMA, str(path))
    return frame, stamps


def _rate(numerator: int, denominator: int) -> float | None:
    """A share rounded for the manifest; None over an empty denominator."""
    if not denominator:
        return None
    return round(numerator / denominator, RATE_DECIMALS)


def unlabeled_figures() -> AccuracyFigures:
    """The figures block when no labeled sample exists: every figure null."""
    return {
        "labeled": False,
        "drawn": None,
        "rejected": None,
        "n": None,
        "seed": None,
        "drawn_at": None,
        "sentiment_agreement": None,
        "target_agreement": None,
        "joint_agreement": None,
        "by_class": None,
    }


def score_sample(
    sample: pl.DataFrame, *, seed: int | None, drawn_at: str | None
) -> AccuracyFigures:
    """
    Score the classifier against the owner's verdicts.

    Rejected rows are excluded before anything is counted. Sentiment
    agreement is the label match; target agreement is the owner naming
    the attributed player (none and other both count against); joint is
    both on the same row. Per class: precision and recall of the label,
    and toward_precision, labeled so and about the attributed player,
    of the predicted, which is the figure a negative rate rests on.

    Args:
        sample: Frame under ACCURACY_SAMPLE_SCHEMA.
        seed: The draw's seed, for the block.
        drawn_at: The draw's date, for the block.

    Returns:
        The AccuracyFigures block, labeled=True.
    """
    scored = sample.filter(pl.col("reject").is_null())
    n = scored.height
    sentiment_ok = pl.col("sentiment") == pl.col("label_sentiment")
    target_ok = pl.col("attributed_player") == pl.col("label_target")

    by_class: dict[str, ClassAgreement] = {}
    for label in SENTIMENT_VERDICTS:
        predicted = scored.filter(pl.col("sentiment") == label)
        labeled = scored.filter(pl.col("label_sentiment") == label)
        agreed = pl.col("label_sentiment") == label
        by_class[label] = {
            "predicted": predicted.height,
            "labeled": labeled.height,
            "precision": _rate(predicted.filter(agreed).height, predicted.height),
            "recall": _rate(
                labeled.filter(pl.col("sentiment") == label).height, labeled.height
            ),
            "toward_precision": _rate(
                predicted.filter(agreed & target_ok).height, predicted.height
            ),
        }
    return {
        "labeled": True,
        "drawn": sample.height,
        "rejected": sample.height - n,
        "n": n,
        "seed": seed,
        "drawn_at": drawn_at,
        "sentiment_agreement": _rate(scored.filter(sentiment_ok).height, n),
        "target_agreement": _rate(scored.filter(target_ok).height, n),
        "joint_agreement": _rate(scored.filter(sentiment_ok & target_ok).height, n),
        "by_class": by_class,
    }


def load_accuracy_sample(
    path: Path | None, fact_stamps: Mapping[str, str | None]
) -> AccuracyFigures:
    """
    Score the labeled sample on disk, or report that there is none.

    The sample's players-config stamp is drift-checked (the target
    options were built under it) and its classifier identity is compared
    with the fact's: a figure describes the classifier it was labeled
    against, so a mismatch is warned, never hidden.

    Args:
        path: Path to accuracy_sample.parquet, or None.
        fact_stamps: The fact's classifier_sentiment stamps.

    Returns:
        The AccuracyFigures block; unlabeled when the file is absent.

    Raises:
        ValueError: If the file does not conform to ACCURACY_SAMPLE_SCHEMA.
    """
    if path is None or not path.exists():
        logger.warning(
            f"no accuracy sample at {path}: the manifest carries no accuracy figure"
        )
        return unlabeled_figures()

    metadata = pl.read_parquet_metadata(path)
    check_config_stamps(
        path,
        metadata,
        "accuracy_sample",
        subject="accuracy sample",
        remedy="the target options were offered under another alias map; the "
        "figure stands for the config it was labeled under",
        log=logger,
    )
    for key in SENTIMENT_STAMP_KEYS:
        if metadata.get(key) != fact_stamps.get(key):
            logger.warning(
                f"{path}: {key} is {metadata.get(key)!r} but the fact carries "
                f"{fact_stamps.get(key)!r}; the accuracy figure describes the "
                f"classifier it was labeled against, not this fact - redraw and "
                f"relabel"
            )

    sample = pl.read_parquet(path)
    validate_schema(sample, ACCURACY_SAMPLE_SCHEMA, str(path))
    seed = metadata.get("sample_seed")
    figures = score_sample(
        sample,
        seed=int(seed) if seed is not None else None,
        drawn_at=metadata.get("drawn_at"),
    )
    log_figures(figures)
    return figures


def log_figures(figures: AccuracyFigures) -> None:
    """
    Log a labeled figures block in one readable pass.

    Args:
        figures: A labeled AccuracyFigures block.
    """

    def pct(value: float | None) -> str:
        return "n/a" if value is None else f"{value:.1%}"

    logger.info(
        f"accuracy sample: n={figures['n']:,} scored of {figures['drawn']:,} drawn "
        f"({figures['rejected']:,} rejected); sentiment "
        f"{pct(figures['sentiment_agreement'])}, target "
        f"{pct(figures['target_agreement'])}, joint {pct(figures['joint_agreement'])}"
    )
    for label, block in (figures["by_class"] or {}).items():
        logger.info(
            f"  {label}: predicted {block['predicted']:,} / labeled "
            f"{block['labeled']:,}; precision {pct(block['precision'])}, recall "
            f"{pct(block['recall'])}, toward {pct(block['toward_precision'])}"
        )
