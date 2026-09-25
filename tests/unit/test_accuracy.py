"""Tests for pipeline.accuracy (the draw, the workbook round trip, the figures)."""

import logging
from pathlib import Path

import polars as pl
import pytest
from openpyxl import load_workbook

from pipeline.accuracy import (
    CLASSIFIER_SHEET,
    LABEL_COLUMNS,
    LABELS_SHEET,
    LISTS_SHEET,
    META_SHEET,
    REJECT_REASONS,
    SAMPLE_COLUMNS,
    SENTIMENT_VERDICTS,
    TARGET_NONE,
    TARGET_OTHER,
    draw_sample,
    load_accuracy_sample,
    read_workbook,
    score_sample,
    target_options,
    unlabeled_figures,
    write_workbook,
)
from pipeline.schemas import ACCURACY_SAMPLE_SCHEMA, AccuracyFigures, ClassAgreement
from utils.constants import ACCURACY_CLASS_MIN_N
from utils.player_config import load_player_config_version

_FACT_SCHEMA = pl.Schema(
    {
        "comment_id": pl.String,
        "link_id": pl.String,
        "body": pl.String,
        "mentioned_players": pl.List(pl.String),
        "sentiment": pl.String,
        "confidence": pl.Float64,
        "sentiment_player": pl.String,
        "attributed_player": pl.String,
    }
)

STAMPS = {
    "players_config_version": "4.6",
    "classifier_sentiment_model": "claude-haiku-4-5-20251001",
    "classifier_sentiment_prompt_version": "v2-production+s-hint",
    "sample_seed": "11",
    "sample_n": "3",
    "drawn_at": "2026-09-24",
}


def _fact(rows: list[dict]) -> pl.DataFrame:
    """A fact-shaped frame with stable defaults per row."""
    defaults = {
        "link_id": "t3_post1",
        "body": "a comment",
        "mentioned_players": ["LeBron James"],
        "sentiment": "neg",
        "confidence": 0.95,
        "sentiment_player": "LeBron James",
        "attributed_player": "LeBron James",
    }
    return pl.DataFrame([{**defaults, **row} for row in rows], schema=_FACT_SCHEMA)


def _sample() -> pl.DataFrame:
    """Three drawn rows: a single mention, a double mention, a neutral."""
    return pl.DataFrame(
        {
            "comment_id": ["c1", "c2", "c3"],
            "permalink": [
                "https://www.reddit.com/r/nba/comments/post1/comment/c1/",
                "https://www.reddit.com/r/nba/comments/post1/comment/c2/",
                "https://www.reddit.com/r/nba/comments/post2/comment/c3/",
            ],
            "post_title": ["Game thread", "Game thread", None],
            "body": ["LeBron washed", "Luka > LeBron", "Jokic played"],
            "mentioned_players": [
                ["LeBron James"],
                ["Luka Doncic", "LeBron James"],
                ["Nikola Jokic"],
            ],
            "sentiment": ["neg", "neg", "neu"],
            "confidence": [0.95, 0.9, 0.5],
            "sentiment_player": ["LeBron James", "LeBron James", None],
            "attributed_player": ["LeBron James", "LeBron James", "Nikola Jokic"],
        },
        schema=pl.Schema(
            {
                "comment_id": pl.String,
                "permalink": pl.String,
                "post_title": pl.String,
                "body": pl.String,
                "mentioned_players": pl.List(pl.String),
                "sentiment": pl.String,
                "confidence": pl.Float64,
                "sentiment_player": pl.String,
                "attributed_player": pl.String,
            }
        ),
    )


def _fill(path: Path, verdicts: dict[str, tuple]) -> None:
    """Fill the labels sheet: comment_id -> (sentiment, target, reject, note)."""
    workbook = load_workbook(path)
    sheet = workbook[LABELS_SHEET]
    for row in sheet.iter_rows(min_row=2):
        comment_id = row[0].value
        if comment_id in verdicts:
            for cell, value in zip(row[4:8], verdicts[comment_id]):
                cell.value = value
    workbook.save(path)


COMPLETE = {
    "c1": ("neg", "LeBron James", None, None),
    "c2": ("neg", "Luka Doncic", None, "the target is Luka"),
    "c3": ("neu", TARGET_NONE, None, None),
}


@pytest.fixture
def workbook(tmp_path) -> Path:
    """The labeling workbook for _sample, written and unfilled."""
    path = tmp_path / "accuracy_sample.xlsx"
    write_workbook(_sample(), path, STAMPS)
    return path


class TestDrawSample:
    """Tests for draw_sample."""

    def _population(self, size: int) -> pl.DataFrame:
        return _fact(
            [
                {"comment_id": f"c{i:02d}", "link_id": f"t3_p{i % 3}"}
                for i in range(size)
            ]
        )

    def test_is_deterministic_under_the_seed(self):
        """The same seed draws the same ids in the same order; another seed differs."""
        population = self._population(40)

        first = draw_sample(population, None, n=10, seed=3)
        second = draw_sample(population, None, n=10, seed=3)
        other = draw_sample(population, None, n=10, seed=4)

        assert first["comment_id"].to_list() == second["comment_id"].to_list()
        assert first["comment_id"].to_list() != other["comment_id"].to_list()
        assert first.columns == list(SAMPLE_COLUMNS)

    def test_population_is_attributed_rows_with_a_real_label(self):
        """Unattributed and error rows are outside the draw."""
        population = _fact(
            [
                {"comment_id": "keep1"},
                {"comment_id": "keep2", "sentiment": "pos"},
                {"comment_id": "drop1", "attributed_player": None},
                {"comment_id": "drop2", "sentiment": "error"},
            ]
        )

        sample = draw_sample(population, None, n=2, seed=1)

        assert sorted(sample["comment_id"].to_list()) == ["keep1", "keep2"]

    def test_raises_when_the_population_is_short(self):
        """Asking for more rows than the population holds is an error, not a smaller draw."""
        with pytest.raises(ValueError, match="fewer than the 5 requested"):
            draw_sample(self._population(3), None, n=5, seed=1)

    def test_permalink_from_post_and_comment_ids(self):
        """The permalink strips the t3_ prefix and points at the comment."""
        population = _fact([{"comment_id": "abc", "link_id": "t3_xyz"}])

        sample = draw_sample(population, None, n=1, seed=1)

        assert sample["permalink"][0] == (
            "https://www.reddit.com/r/nba/comments/xyz/comment/abc/"
        )

    def test_titles_join_from_the_bridge(self):
        """With a bridge each row carries its post's title; an unbridged post is null."""
        population = _fact(
            [
                {"comment_id": "a", "link_id": "t3_known"},
                {"comment_id": "b", "link_id": "t3_unknown"},
            ]
        )
        posts = pl.DataFrame({"post_id": ["t3_known"], "title": ["Game thread"]})

        sample = draw_sample(population, posts, n=2, seed=1).sort("comment_id")

        assert sample["post_title"].to_list() == ["Game thread", None]

    def test_without_a_bridge_titles_are_null(self):
        """No bridge given: the column exists, every title null."""
        sample = draw_sample(self._population(4), None, n=2, seed=1)

        assert sample["post_title"].dtype == pl.String
        assert sample["post_title"].null_count() == 2

    def test_warns_when_a_predicted_class_is_thin(self, caplog):
        """A class under the floor in the draw is flagged before labeling starts."""
        population = self._population(ACCURACY_CLASS_MIN_N + 5)

        with caplog.at_level(logging.WARNING, logger="pipeline.accuracy"):
            draw_sample(population, None, n=ACCURACY_CLASS_MIN_N, seed=1)

        assert "fall under" in caplog.text
        assert "'neu', 'pos'" in caplog.text


class TestTargetOptions:
    """Tests for target_options."""

    def test_mentions_then_none_then_other(self):
        """The list is the row's mentions in order, then the two escapes."""
        assert target_options(["Luka Doncic", "LeBron James"]) == (
            "Luka Doncic",
            "LeBron James",
            TARGET_NONE,
            TARGET_OTHER,
        )


class TestWriteWorkbook:
    """Tests for write_workbook: what the labeler sees and what they don't."""

    def test_labels_sheet_carries_no_prediction(self, workbook):
        """The visible sheet is the comment plus empty verdict columns."""
        sheet = load_workbook(workbook)[LABELS_SHEET]

        header = [cell.value for cell in sheet[1]]
        rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]

        assert header == list(LABEL_COLUMNS)
        assert rows[0][:4] == ["c1", "open", "Game thread", "LeBron washed"]
        assert rows[0][4:] == [None, None, None, None]
        assert sheet["B2"].hyperlink.target == _sample()["permalink"][0]

    def test_classifier_lists_and_meta_sheets_are_hidden(self, workbook):
        """The join keys and the draw's stamps ride along out of sight."""
        book = load_workbook(workbook)

        assert book[LABELS_SHEET].sheet_state == "visible"
        for name in (CLASSIFIER_SHEET, LISTS_SHEET, META_SHEET):
            assert book[name].sheet_state == "hidden"

    def test_verdict_columns_are_list_validated(self, workbook):
        """Sentiment and reject take one list each over the whole column."""
        sheet = load_workbook(workbook)[LABELS_SHEET]
        validations = {
            str(v.sqref): v.formula1 for v in sheet.data_validations.dataValidation
        }

        assert validations["E2:E4"] == '"' + ",".join(SENTIMENT_VERDICTS) + '"'
        assert validations["G2:G4"] == '"' + ",".join(REJECT_REASONS) + '"'

    def test_a_typed_value_outside_the_list_is_refused_in_the_cell(self, workbook):
        """Every dropdown raises a stop alert, so a typo is caught while labeling."""
        sheet = load_workbook(workbook)[LABELS_SHEET]

        for validation in sheet.data_validations.dataValidation:
            assert validation.showErrorMessage is True
            assert validation.errorStyle == "stop"
            assert validation.allow_blank is True

    def test_target_list_is_per_row(self, workbook):
        """Each row's target dropdown is its own mentions plus none and other."""
        book = load_workbook(workbook)
        lists = [
            [cell.value for cell in row if cell.value is not None]
            for row in book[LISTS_SHEET].iter_rows()
        ]
        by_cell = {}
        for validation in book[LABELS_SHEET].data_validations.dataValidation:
            for cell_range in str(validation.sqref).split():
                by_cell[cell_range] = validation.formula1

        assert lists == [
            ["LeBron James", TARGET_NONE, TARGET_OTHER],
            ["Luka Doncic", "LeBron James", TARGET_NONE, TARGET_OTHER],
            ["Nikola Jokic", TARGET_NONE, TARGET_OTHER],
        ]
        assert by_cell["F2"] == f"{LISTS_SHEET}!$A$1:$C$1"
        assert by_cell["F3"] == f"{LISTS_SHEET}!$A$2:$D$2"
        assert by_cell["F4"] == f"{LISTS_SHEET}!$A$3:$C$3"

    def test_rows_sharing_a_list_share_a_validation(self, tmp_path):
        """Distinct lists, not rows, drive the lists sheet."""
        path = tmp_path / "book.xlsx"
        sample = pl.concat(
            [
                _sample(),
                _sample().with_columns(pl.col("comment_id").str.replace("c", "d")),
            ]
        )

        write_workbook(sample, path, STAMPS)

        book = load_workbook(path)
        assert book[LISTS_SHEET].max_row == 3
        assert len(book[LABELS_SHEET].data_validations.dataValidation) == 5


class TestReadWorkbook:
    """Tests for read_workbook: the round trip and every refusal."""

    def test_round_trip(self, workbook):
        """A complete workbook comes back as the sample in draw order, with its stamps."""
        _fill(workbook, COMPLETE)

        sample, stamps = read_workbook(workbook)

        assert sample.schema == ACCURACY_SAMPLE_SCHEMA
        assert sample["comment_id"].to_list() == ["c1", "c2", "c3"]
        assert sample["label_sentiment"].to_list() == ["neg", "neg", "neu"]
        assert sample["label_target"].to_list() == [
            "LeBron James",
            "Luka Doncic",
            TARGET_NONE,
        ]
        assert sample["note"].to_list() == [None, "the target is Luka", None]
        assert sample["sentiment"].to_list() == ["neg", "neg", "neu"]
        assert sample["sentiment_player"].to_list() == [
            "LeBron James",
            "LeBron James",
            None,
        ]
        assert sample["attributed_player"][2] == "Nikola Jokic"
        assert sample["confidence"].to_list() == [0.95, 0.9, 0.5]
        assert stamps == STAMPS

    def test_rejected_row_carries_the_reason_and_no_labels(self, workbook):
        """A reject needs no verdicts; whatever was typed beside it is dropped."""
        _fill(
            workbook,
            {**COMPLETE, "c2": ("pos", "other", "alias_false_positive", "a referee")},
        )

        sample, _ = read_workbook(workbook)

        assert sample.row(1, named=True) == {
            "comment_id": "c2",
            "sentiment": "neg",
            "confidence": 0.9,
            "sentiment_player": "LeBron James",
            "attributed_player": "LeBron James",
            "label_sentiment": None,
            "label_target": None,
            "reject": "alias_false_positive",
            "note": "a referee",
        }

    def test_sorted_rows_still_read(self, workbook):
        """The join is by comment id, so a sorted sheet is fine."""
        _fill(workbook, COMPLETE)
        book = load_workbook(workbook)
        sheet = book[LABELS_SHEET]
        rows = [[cell.value for cell in row] for row in sheet.iter_rows(min_row=2)]
        for index, row in enumerate(reversed(rows), start=2):
            for column, value in enumerate(row, start=1):
                sheet.cell(row=index, column=column, value=value)
        book.save(workbook)

        sample, _ = read_workbook(workbook)

        assert sample["comment_id"].to_list() == ["c1", "c2", "c3"]

    def test_trailing_empty_rows_are_ignored(self, workbook):
        """A stray blank row under the table is not a problem."""
        _fill(workbook, COMPLETE)
        book = load_workbook(workbook)
        book[LABELS_SHEET].cell(row=6, column=8, value=None)
        book.save(workbook)

        sample, _ = read_workbook(workbook)

        assert sample.height == 3

    @pytest.mark.parametrize(
        "verdicts,match",
        [
            ({"c1": (None, "LeBron James", None, None)}, r"row 2: sentiment None"),
            (
                {"c1": ("positive", "LeBron James", None, None)},
                r"row 2: sentiment 'positive'",
            ),
            ({"c1": ("neg", None, None, None)}, r"row 2: target None is not one of"),
            (
                {"c2": ("neg", "Nikola Jokic", None, None)},
                r"row 3: target 'Nikola Jokic'",
            ),
            ({"c3": (None, None, "spam", None)}, r"row 4: reject 'spam'"),
        ],
        ids=[
            "blank-sentiment",
            "off-list-sentiment",
            "blank-target",
            "off-row-target",
            "unknown-reject",
        ],
    )
    def test_refuses_an_off_list_verdict_naming_the_row(
        self, workbook, verdicts, match
    ):
        """Every verdict must come from its list; the message names the row."""
        _fill(workbook, {**COMPLETE, **verdicts})

        with pytest.raises(ValueError, match=match):
            read_workbook(workbook)

    def test_refuses_a_missing_row(self, workbook):
        """Every drawn id must be labeled; a deleted row is named."""
        _fill(workbook, COMPLETE)
        book = load_workbook(workbook)
        book[LABELS_SHEET].delete_rows(3)
        book.save(workbook)

        with pytest.raises(ValueError, match=r"1 drawn row\(s\) are absent.*c2"):
            read_workbook(workbook)

    def test_refuses_a_duplicate_and_an_unknown_id(self, workbook):
        """A copied row and a row outside the draw are both named."""
        _fill(workbook, COMPLETE)
        book = load_workbook(workbook)
        sheet = book[LABELS_SHEET]
        sheet.append(["c1", None, None, None, "neg", "LeBron James", None, None])
        sheet.append(["zz", None, None, None, "neg", TARGET_NONE, None, None])
        book.save(workbook)

        with pytest.raises(ValueError) as excinfo:
            read_workbook(workbook)

        assert "row 5: c1 appears twice" in str(excinfo.value)
        assert "row 6: zz is not in the draw" in str(excinfo.value)

    def test_reports_every_problem_at_once(self, workbook):
        """One pass names all the rows, so the fix is one edit session."""
        _fill(workbook, {"c2": COMPLETE["c2"]})

        with pytest.raises(ValueError) as excinfo:
            read_workbook(workbook)

        assert "2 problem(s)" in str(excinfo.value)
        assert "row 2: unlabeled" in str(excinfo.value)
        assert "row 4: unlabeled" in str(excinfo.value)

    def test_refuses_a_changed_header(self, workbook):
        """The columns are the contract; a renamed header is not read around."""
        book = load_workbook(workbook)
        book[LABELS_SHEET]["E1"] = "Sentiment"
        book.save(workbook)

        with pytest.raises(ValueError, match="header was changed"):
            read_workbook(workbook)

    def test_refuses_a_missing_sheet(self, workbook):
        """Without the hidden join sheets the labels cannot be scored."""
        book = load_workbook(workbook)
        del book[CLASSIFIER_SHEET]
        book.save(workbook)

        with pytest.raises(ValueError, match="sheet 'classifier' is missing"):
            read_workbook(workbook)


def _labeled(rows: list[dict]) -> pl.DataFrame:
    """A labeled sample with per-row defaults: LeBron, predicted neg, labeled neg toward him."""
    defaults = {
        "sentiment": "neg",
        "confidence": 0.9,
        "sentiment_player": "LeBron James",
        "attributed_player": "LeBron James",
        "label_sentiment": "neg",
        "label_target": "LeBron James",
        "reject": None,
        "note": None,
    }
    return pl.DataFrame(
        [{"comment_id": f"c{i}", **defaults, **row} for i, row in enumerate(rows)],
        schema=ACCURACY_SAMPLE_SCHEMA,
    )


CONFUSION = _labeled(
    [
        {},  # neg right, toward right
        {"label_sentiment": "neu"},  # neg wrong (neu), toward right
        {"label_target": TARGET_NONE},  # neg right, toward wrong
        {"sentiment": "pos", "label_sentiment": "pos", "label_target": TARGET_OTHER},
        {"sentiment": "neu", "label_sentiment": "neu"},
        {"label_sentiment": None, "label_target": None, "reject": "unreadable"},
    ]
)


class TestScoreSample:
    """Tests for score_sample."""

    def test_headline_figures(self):
        """Agreement is counted over scored rows; the reject is out of every denominator."""
        figures = score_sample(CONFUSION, seed=11, drawn_at="2026-09-24")

        assert figures["labeled"] is True
        assert (figures["drawn"], figures["rejected"], figures["n"]) == (6, 1, 5)
        assert (figures["seed"], figures["drawn_at"]) == (11, "2026-09-24")
        assert figures["sentiment_agreement"] == 0.8
        assert figures["target_agreement"] == 0.6
        assert figures["joint_agreement"] == 0.4

    def test_per_class_figures(self):
        """Precision, recall and toward-precision per label."""
        by_class = score_sample(CONFUSION, seed=None, drawn_at=None)["by_class"]

        assert by_class["neg"] == {
            "predicted": 3,
            "labeled": 2,
            "precision": 0.6667,
            "recall": 1.0,
            "toward_precision": 0.3333,
        }
        assert by_class["pos"] == {
            "predicted": 1,
            "labeled": 1,
            "precision": 1.0,
            "recall": 1.0,
            "toward_precision": 0.0,
        }
        assert by_class["neu"] == {
            "predicted": 1,
            "labeled": 2,
            "precision": 1.0,
            "recall": 0.5,
            "toward_precision": 1.0,
        }

    def test_empty_denominators_are_null(self):
        """Nothing scored, nothing claimed: every rate null, counts zero."""
        figures = score_sample(
            _labeled(
                [
                    {
                        "label_sentiment": None,
                        "label_target": None,
                        "reject": "unreadable",
                    }
                ]
            ),
            seed=1,
            drawn_at=None,
        )

        assert figures["n"] == 0
        assert figures["sentiment_agreement"] is None
        assert figures["by_class"]["neg"]["precision"] is None

    def test_blocks_match_the_contract(self):
        """Both figures blocks carry exactly the contract's fields, in order."""
        figures = score_sample(CONFUSION, seed=1, drawn_at=None)

        assert list(figures) == list(AccuracyFigures.__annotations__)
        assert list(unlabeled_figures()) == list(AccuracyFigures.__annotations__)
        assert list(figures["by_class"]["neg"]) == list(ClassAgreement.__annotations__)


class TestLoadAccuracySample:
    """Tests for load_accuracy_sample: the file on disk, or none."""

    FACT_STAMPS = {
        "classifier_sentiment_model": "claude-haiku-4-5-20251001",
        "classifier_sentiment_prompt_version": "v2-production+s-hint",
    }

    def _write(self, path: Path, **overrides) -> Path:
        metadata = {
            "players_config_version": load_player_config_version(),
            **self.FACT_STAMPS,
            "sample_seed": "11",
            "drawn_at": "2026-09-24",
            **overrides,
        }
        CONFUSION.write_parquet(
            path, metadata={k: v for k, v in metadata.items() if v is not None}
        )
        return path

    def test_absent_file_is_unlabeled(self, tmp_path, caplog):
        """No file: the unlabeled block, and a warning that says so."""
        with caplog.at_level(logging.WARNING, logger="pipeline.accuracy"):
            figures = load_accuracy_sample(
                tmp_path / "missing.parquet", self.FACT_STAMPS
            )

        assert figures == unlabeled_figures()
        assert "no accuracy figure" in caplog.text

    def test_none_path_is_unlabeled(self):
        """No path at all reads the same as a missing file."""
        assert load_accuracy_sample(None, self.FACT_STAMPS) == unlabeled_figures()

    def test_scores_the_file_with_its_stamps(self, tmp_path, caplog):
        """The figures come from the file, seed and date from its metadata, no warning."""
        path = self._write(tmp_path / "accuracy_sample.parquet")

        with caplog.at_level(logging.WARNING, logger="pipeline.accuracy"):
            figures = load_accuracy_sample(path, self.FACT_STAMPS)

        assert figures["labeled"] is True
        assert figures["n"] == 5
        assert figures["seed"] == 11
        assert figures["drawn_at"] == "2026-09-24"
        assert caplog.text == ""

    def test_classifier_mismatch_warns(self, tmp_path, caplog):
        """A sample labeled against another classifier is scored, but said so."""
        path = self._write(
            tmp_path / "accuracy_sample.parquet",
            classifier_sentiment_prompt_version="v1",
        )

        with caplog.at_level(logging.WARNING, logger="pipeline.accuracy"):
            figures = load_accuracy_sample(path, self.FACT_STAMPS)

        assert figures["labeled"] is True
        assert "describes the classifier it was labeled against" in caplog.text
        assert "'v1'" in caplog.text

    def test_config_drift_warns(self, tmp_path, caplog):
        """The players-config stamp is drift-checked like any config-derived file."""
        path = self._write(
            tmp_path / "accuracy_sample.parquet", players_config_version="0.1"
        )

        with caplog.at_level(logging.WARNING, logger="pipeline.accuracy"):
            load_accuracy_sample(path, self.FACT_STAMPS)

        assert "players_config_version drift" in caplog.text

    def test_wrong_shape_raises(self, tmp_path):
        """A parquet that is not the sample is refused, not scored."""
        path = tmp_path / "accuracy_sample.parquet"
        pl.DataFrame({"comment_id": ["c1"]}).write_parquet(path)

        with pytest.raises(ValueError):
            load_accuracy_sample(path, self.FACT_STAMPS)
